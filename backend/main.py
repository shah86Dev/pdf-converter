from io import BytesIO
from pathlib import Path
import re

import pdfplumber
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook

app = FastAPI(title="Tabledrop PDF Converter")
MAX_BYTES = 25 * 1024 * 1024

@app.get("/health")
def health():
    return {"status": "ok"}

@app.post("/convert")
async def convert(file: UploadFile = File(...)):
    name = file.filename or "document.pdf"
    if not name.lower().endswith(".pdf") and file.content_type != "application/pdf":
        raise HTTPException(status_code=415, detail="Only PDF files are supported")
    data = await file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise HTTPException(status_code=413, detail="File exceeds the 25 MB limit")
    if not data or not data.startswith(b"%PDF"):
        raise HTTPException(status_code=400, detail="The uploaded file is not a valid PDF")

    workbook = Workbook()
    workbook.remove(workbook.active)
    try:
        with pdfplumber.open(BytesIO(data)) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                tables = page.extract_tables() or []
                if tables:
                    for table_number, table in enumerate(tables, start=1):
                        sheet = workbook.create_sheet(f"Page {page_number} - Table {table_number}"[:31])
                        for row in table:
                            sheet.append([cell or "" for cell in row])
                else:
                    text = page.extract_text() or ""
                    if text.strip():
                        sheet = workbook.create_sheet(f"Page {page_number}"[:31])
                        for line in text.splitlines():
                            sheet.append([line])
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Could not read this PDF") from exc

    if not workbook.worksheets:
        sheet = workbook.create_sheet("Extracted data")
        sheet.append(["No readable tables or text were found in this PDF."])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 48
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    stem = re.sub(r"[^A-Za-z0-9_-]+", "-", Path(name).stem).strip("-") or "converted-document"
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{stem}.xlsx"'})
